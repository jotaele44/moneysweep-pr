#!/usr/bin/env python3
"""Deterministically audit Puerto Rico Planning Board macro workbooks.

The harness is intentionally source-preserving: it hashes the original XLSX bytes,
inventories ZIP members and worksheet visibility, keeps formula and cached-value
manifests separate, and searches every worksheet without assuming row 1 is a
header. It does not infer sector allocations from GDP, net income, payroll,
employment, exports, facilities, incentives, or historical shares.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path
from typing import Any
from zipfile import ZipFile

import openpyxl

PATTERNS = [
    re.compile(value, re.IGNORECASE)
    for value in [
        r"inversi[oó]n\s+directa",
        r"direct\s+investment",
        r"ganancias",
        r"profits?",
        r"dividend",
        r"sector",
        r"industr",
        r"manufact",
        r"farmac",
        r"pharma",
        r"qu[ií]mic",
        r"chemical",
        r"resto\s+del\s+mundo",
        r"rest\s+of\s+world",
    ]
]


def sha256_path(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def zip_members(path: Path) -> list[dict[str, Any]]:
    with ZipFile(path) as archive:
        return [
            {
                "path": info.filename,
                "uncompressed_size": info.file_size,
                "compressed_size": info.compress_size,
            }
            for info in archive.infolist()
        ]


def _cell_text(value: Any) -> str:
    return "" if value is None else str(value)


def semantic_hits(path: Path) -> list[dict[str, Any]]:
    workbook = openpyxl.load_workbook(path, data_only=False, read_only=True, keep_links=True)
    hits: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            texts = [_cell_text(value) for value in row]
            joined = " | ".join(texts)
            if any(pattern.search(joined) for pattern in PATTERNS):
                hits.append(
                    {
                        "sheet": sheet.title,
                        "sheet_state": sheet.sheet_state,
                        "row_number": row_number,
                        "row_values": texts,
                    }
                )
    workbook.close()
    return hits


def worksheet_inventory(path: Path) -> list[dict[str, Any]]:
    formula_book = openpyxl.load_workbook(path, data_only=False, read_only=False, keep_links=True)
    value_book = openpyxl.load_workbook(path, data_only=True, read_only=False, keep_links=True)
    inventory: list[dict[str, Any]] = []
    for name in formula_book.sheetnames:
        formula_sheet = formula_book[name]
        value_sheet = value_book[name]
        formula_cells = 0
        cached_nonempty = 0
        for row in formula_sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_cells += 1
        for row in value_sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cached_nonempty += 1
        inventory.append(
            {
                "sheet": name,
                "state": formula_sheet.sheet_state,
                "max_row": formula_sheet.max_row,
                "max_column": formula_sheet.max_column,
                "merged_ranges": [str(value) for value in formula_sheet.merged_cells.ranges],
                "formula_cell_count": formula_cells,
                "cached_nonempty_cell_count": cached_nonempty,
            }
        )
    formula_book.close()
    value_book.close()
    return inventory


def audit_workbook(path: Path) -> dict[str, Any]:
    if path.suffix.lower() != ".xlsx":
        raise ValueError(f"expected .xlsx workbook: {path}")
    if not path.is_file():
        raise FileNotFoundError(path)
    return {
        "path": str(path),
        "size_bytes": path.stat().st_size,
        "sha256": sha256_path(path),
        "zip_members": zip_members(path),
        "worksheets": worksheet_inventory(path),
        "semantic_hits": semantic_hits(path),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbooks", nargs="+", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()

    payload = {
        "audit_version": "1.0",
        "workbooks": [audit_workbook(path) for path in args.workbooks],
        "hard_gates": [
            "WHOLE_ROW_ONLY",
            "NO_ROW1_HEADER_ASSUMPTION",
            "FORMULA_AND_CACHED_VALUE_MANIFESTATIONS_SEPARATE",
            "INDUSTRY_GDP_IS_NOT_DIRECT_INVESTMENT_PROFIT",
            "UNKNOWN_IS_NOT_ZERO",
        ],
    }
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"wrote workbook audit -> {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
