from __future__ import annotations

import json
from pathlib import Path

import openpyxl

from scripts.audit_jp_macro_workbooks import audit_workbook, main


def _fixture(path: Path) -> None:
    workbook = openpyxl.Workbook()
    visible = workbook.active
    visible.title = "Current"
    visible.append(["Puerto Rico Planning Board"])
    visible.append(["Sector", "FY2025", "Notes"])
    visible.append(["Manufacturing", 100.0, "Direct investment profits candidate"])
    visible["D3"] = "=B3*2"

    hidden = workbook.create_sheet("Auxiliary")
    hidden.sheet_state = "hidden"
    hidden.append(["Ganancias de inversión directa por sector", 41_664.1])

    very_hidden = workbook.create_sheet("Internal")
    very_hidden.sheet_state = "veryHidden"
    very_hidden.append(["rest of world", "diagnostic"])
    workbook.save(path)
    workbook.close()


def test_audit_preserves_bytes_structure_and_hidden_sheets(tmp_path: Path) -> None:
    path = tmp_path / "jp.xlsx"
    _fixture(path)

    result = audit_workbook(path)

    assert result["size_bytes"] > 0
    assert len(result["sha256"]) == 64
    assert any(member["path"] == "xl/workbook.xml" for member in result["zip_members"])
    states = {row["sheet"]: row["state"] for row in result["worksheets"]}
    assert states == {"Current": "visible", "Auxiliary": "hidden", "Internal": "veryHidden"}
    current = next(row for row in result["worksheets"] if row["sheet"] == "Current")
    assert current["formula_cell_count"] == 1


def test_semantic_search_returns_whole_rows_without_header_assumption(tmp_path: Path) -> None:
    path = tmp_path / "jp.xlsx"
    _fixture(path)

    result = audit_workbook(path)
    hits = result["semantic_hits"]

    assert any(hit["sheet"] == "Auxiliary" and hit["row_number"] == 1 for hit in hits)
    assert any("Manufacturing" in hit["row_values"] for hit in hits)
    assert any(hit["sheet_state"] == "veryHidden" for hit in hits)


def test_cli_writes_deterministic_shape(tmp_path: Path, monkeypatch) -> None:
    path = tmp_path / "jp.xlsx"
    out = tmp_path / "audit.json"
    _fixture(path)
    monkeypatch.setattr("sys.argv", ["audit_jp_macro_workbooks.py", str(path), "--out", str(out)])

    assert main() == 0
    payload = json.loads(out.read_text(encoding="utf-8"))
    assert payload["audit_version"] == "1.0"
    assert len(payload["workbooks"]) == 1
    assert "WHOLE_ROW_ONLY" in payload["hard_gates"]
